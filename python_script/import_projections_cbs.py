"""
Importe les projections CBS Sports (aide au choix des poolers, préparation de repêchage/agents
libres) dans player_projections. Usage ponctuel (une fois par saison), pas un pipeline récurrent.

Remplace l'import ESPN (abandonné, voir SUIVI_PROJET.md session 2026-09-13) — le tableau ESPN
collé depuis le site n'avait aucune garantie d'alignement fiable entre la colonne nom et les
colonnes de stats (deux widgets scrollés indépendamment sur la page). Le fichier CBS n'a pas ce
problème : une seule ligne par joueur, directement collée depuis cbssports.com, colonne "Player"
au format "Nom\xa0POS\xa0\xa0ÉQUIPE" (espaces insécables) — voir CBS_TEAM_ALIASES pour les codes
d'équipe qui diffèrent des codes LNH standards (LV, MON, CLB, TB, WAS, LA, NJ, SJ).

Seuls Nom/Équipe/Points (patineurs, colonne 'p' — déjà buts+passes) ou Victoires (gardiens,
colonne 'w') sont importés, arrondis à l'entier — le reste (fpts, fppg, +/-, TOI, sog...) reste
dans le fichier Excel de David, pas stocké dans l'app.

Un contrôle de cohérence compare chaque projection de points patineur à la projection NHL.com
déjà en base (si présente) — un grand écart (>40 pts) signale un jumelage ou un nom suspect.

Usage:
    python import_projections_cbs.py <fichier.xlsx>              # dry-run
    python import_projections_cbs.py <fichier.xlsx> --apply       # écrit réellement
"""

import argparse
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
from dotenv import load_dotenv
from supabase import create_client

from projections_common import build_player_lookup, match_player, get_active_season

load_dotenv()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SUPABASE_URL = os.getenv('SUPABASE_URL')
SUPABASE_KEY = os.getenv('SUPABASE_SERVICE_KEY')

SHEETS = {'Attaquants': False, 'Defenseurs': False, 'Gardiens': True}  # valeur = is_goalie

NHL_TEAM_CODES = {
    'ANA', 'BOS', 'BUF', 'CAR', 'CBJ', 'CGY', 'CHI', 'COL', 'DAL', 'DET',
    'EDM', 'FLA', 'LAK', 'MIN', 'MTL', 'NJD', 'NSH', 'NYI', 'NYR', 'OTT',
    'PHI', 'PIT', 'SEA', 'SJS', 'STL', 'TBL', 'TOR', 'UTA', 'VAN', 'VGK',
    'WPG', 'WSH',
}
# CBS abrège certaines équipes différemment des codes LNH standards utilisés dans `teams`.
CBS_TEAM_ALIASES = {
    'LV': 'VGK', 'MON': 'MTL', 'CLB': 'CBJ', 'TB': 'TBL',
    'WAS': 'WSH', 'LA': 'LAK', 'NJ': 'NJD', 'SJ': 'SJS',
}


def normalize_team(code: str) -> str:
    return CBS_TEAM_ALIASES.get(code, code)


def parse_player_cell(raw):
    """'Nathan MacKinnon\xa0C\xa0\xa0COL' -> ('Nathan MacKinnon', 'COL'). None si format inattendu."""
    if not isinstance(raw, str):
        return None, None
    parts = [p for p in raw.split('\xa0') if p]
    if len(parts) < 3:
        return None, None
    return parts[0].strip(), normalize_team(parts[-1].strip())


def parse_sheet(ws, is_goalie: bool):
    """Retourne (records, erreurs). record = {name, team, points|None, wins|None, row}."""
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip().lower() if h else '' for h in rows[0]]
    stat_col = header.index('w') if is_goalie else header.index('p')

    records, errors = [], []
    for i, row in enumerate(rows[1:], start=2):
        if not row or not row[0]:
            continue
        name, team = parse_player_cell(row[0])
        if name is None:
            errors.append(f'Ligne {i} : format de nom inattendu {row[0]!r}')
            continue
        if team not in NHL_TEAM_CODES:
            errors.append(f'Ligne {i} : code équipe {team!r} non reconnu pour {name!r}')
            continue

        raw_val = row[stat_col] if stat_col < len(row) else None
        value = None
        if raw_val not in (None, '', '--'):
            try:
                value = round(float(raw_val))
            except (TypeError, ValueError):
                pass

        records.append({
            'name': name, 'team': team,
            'points': None if is_goalie else value,
            'wins': value if is_goalie else None,
            'row': i,
        })
    return records, errors


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('fichier')
    parser.add_argument('--season')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.fichier, data_only=True)

    skaters, goalies, parse_errors = [], [], []
    for sheet_name, is_goalie in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f'[INFO] Feuille {sheet_name!r} absente du fichier — ignorée.')
            continue
        sheet_records, sheet_errors = parse_sheet(wb[sheet_name], is_goalie)
        print(f'[INFO] {sheet_name} : {len(sheet_records)} joueur(s) reconnu(s), {len(sheet_errors)} ligne(s) non reconnue(s).')
        (goalies if is_goalie else skaters).extend(sheet_records)
        parse_errors.extend(f'{sheet_name} — {e}' for e in sheet_errors)

    print(f'\n[INFO] Total : {len(skaters)} patineur(s), {len(goalies)} gardien(s).')
    if parse_errors:
        print(f'[ATTENTION] {len(parse_errors)} ligne(s) non reconnue(s) :')
        for e in parse_errors:
            print(f'  - {e}')

    db = create_client(SUPABASE_URL, SUPABASE_KEY)
    season = args.season or get_active_season(db)
    print(f'[INFO] Saison ciblée : {season}')

    by_name_team, by_name, by_lastname_team = build_player_lookup(db)

    def resolve(recs):
        matched, unmatched, ambiguous = [], [], []
        for rec in recs:
            pid, note = match_player(rec['name'], rec['team'], by_name_team, by_name, by_lastname_team)
            if pid:
                matched.append((rec, pid, note))
            elif note and 'même nom' in note:
                ambiguous.append((rec, note))
            else:
                unmatched.append((rec, note))
        return matched, unmatched, ambiguous

    def report(label, matched, unmatched, ambiguous):
        print(f'\n[{label}] {len(matched)} jumelé(s), {len(unmatched)} non trouvé(s), {len(ambiguous)} ambigu(s).')
        for rec, note in unmatched:
            print(f'  [NON TROUVÉ] {rec["name"]} ({rec["team"]}) — {note}')
        for rec, note in ambiguous:
            print(f'  [AMBIGU] {rec["name"]} ({rec["team"]}) — {note}')
        approx = [(rec, note) for rec, _, note in matched if note]
        if approx:
            print(f'  {len(approx)} jumelage(s) approximatif(s) (à vérifier) :')
            for rec, note in approx:
                print(f'    [~] {rec["name"]} ({rec["team"]}) — {note}')

    skater_matched, skater_unmatched, skater_ambiguous = resolve(skaters)
    goalie_matched, goalie_unmatched, goalie_ambiguous = resolve(goalies)
    report('PATINEURS', skater_matched, skater_unmatched, skater_ambiguous)
    report('GARDIENS', goalie_matched, goalie_unmatched, goalie_ambiguous)

    # Contrôle de cohérence : compare aux projections NHL.com déjà en base.
    player_ids = [pid for _, pid, _ in skater_matched]
    if player_ids:
        existing = db.table('player_projections').select('player_id, projected_points') \
            .eq('season', season).eq('source', 'nhl_com').in_('player_id', player_ids).execute()
        nhl_points = {row['player_id']: row['projected_points'] for row in existing.data if row['projected_points'] is not None}
        suspects = []
        for rec, pid, _ in skater_matched:
            ref = nhl_points.get(pid)
            if ref is not None and rec['points'] is not None and abs(rec['points'] - ref) > 40:
                suspects.append((rec, ref))
        if suspects:
            print(f'\n[ATTENTION] {len(suspects)} projection(s) CBS très éloignée(s) de NHL.com (>40 pts d\'écart) :')
            for rec, ref in suspects:
                print(f'  [!] {rec["name"]} ({rec["team"]}) — CBS={rec["points"]} pts, NHL.com={ref} pts')

    if not args.apply:
        print('\n[DRY-RUN] Aucune écriture. Relancez avec --apply pour importer les jumelages ci-dessus.')
        return

    all_matched = skater_matched + goalie_matched
    if not all_matched:
        print('\n[INFO] Rien à importer.')
        return

    confirm = input(f'\nImporter {len(all_matched)} projection(s) pour la saison {season} ? (oui/non) ')
    if confirm.strip().lower() != 'oui':
        print('Annulé.')
        return

    rows = []
    for rec, pid, _ in skater_matched:
        rows.append({'player_id': pid, 'season': season, 'source': 'cbs', 'projected_points': rec['points']})
    for rec, pid, _ in goalie_matched:
        rows.append({'player_id': pid, 'season': season, 'source': 'cbs', 'projected_wins': rec['wins']})
    db.table('player_projections').upsert(rows, on_conflict='player_id,season,source').execute()
    print(f'[OK] {len(rows)} projection(s) importée(s)/mise(s) à jour.')


if __name__ == '__main__':
    main()
